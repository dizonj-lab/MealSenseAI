import copy, re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CH = '/tmp/claude-0/-home-user/77ae2d44-c4c4-55cb-8f54-e8e0948aa3c6/scratchpad/report/charts/'
d = docx.Document('reviewed.docx')
P = d.paragraphs
caption_tpl = next(p for p in P if p.style.name == 'Caption' and p.text.startswith('Figure'))
img_tpl = next(p for p in P if p._p.xpath('.//pic:pic'))

def after(p, new_el):
    p._p.addnext(new_el); return new_el

def new_par_after(p, style=None):
    el = OxmlElement('w:p'); p._p.addnext(el)
    np = docx.text.paragraph.Paragraph(el, p._parent)
    if style: np.style = style
    return np

def figure_after(p, png, caption, width=6.2):
    ip = new_par_after(p)
    ip._p.insert(0, copy.deepcopy(img_tpl._p.pPr)) if img_tpl._p.pPr is not None else None
    ip.add_run().add_picture(CH + png, width=Inches(width))
    cp = new_par_after(ip)
    cp._p.insert(0, copy.deepcopy(caption_tpl._p.pPr))
    r = cp.add_run('Figure 0. '); r.bold = True
    src = caption_tpl.runs[0]
    for rr in (r, cp.add_run(caption)):
        rr.font.size = src.font.size; rr.font.color.rgb = src.font.color.rgb if src.font.color and src.font.color.type else None
    return cp

def remove(p):
    p._p.getparent().remove(p._p)

def plain_after(p, text):
    np = new_par_after(p)
    pPr = np._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'E6F2F0'); pPr.append(shd)
    bdr = OxmlElement('w:pBdr'); l = OxmlElement('w:left')
    for k, v in (('w:val', 'single'), ('w:sz', '18'), ('w:space', '6'), ('w:color', '147D73')): l.set(qn(k), v)
    bdr.append(l); pPr.insert(0, bdr)
    sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), '60'); sp.set(qn('w:after'), '160'); pPr.append(sp)
    ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '170'); ind.set(qn('w:right'), '170'); pPr.append(ind)
    r1 = np.add_run('In plain terms: '); r1.bold = True; r1.font.color.rgb = RGBColor(0x14, 0x7D, 0x73)
    np.add_run(text)
    for r in np.runs: r.font.size = Pt(10.5)
    return np

def find(pred):
    return next(p for p in d.paragraphs if pred(p))

def red_para(text):
    return find(lambda p: p.text.strip() == text and any(r.font.color is not None and r.font.color.type is not None and str(r.font.color.rgb) in ('FF0000', 'C00000') for r in p.runs))

reds = [p for p in d.paragraphs if p.text.strip().startswith('(Provide')]
assert len(reds) == 9, len(reds)
r_data, r_train, r_mf, r_cite, r_adv, r_learn, r_sys, r_exp, r_legend = reds

# 1 data preparation diagram
figure_after(r_data, '14_data_prep.png', 'Data preparation: from 310,861 collected images to the three sets used for training, validation and test.'); remove(r_data)
# 2 training diagram
figure_after(r_train, '15_training_flow.png', 'Training, testing and deploying the food classifier.'); remove(r_train)
# 3 evidence for multi-food
figure_after(r_mf, '16_multifood_evidence.png', 'Evidence of multi-food recognition: regions the application found on a test plate, and what the classifier returned for each (E3).'); remove(r_mf)
# 4 citation for the glucose formula
c = new_par_after(r_cite)
c.add_run('Sources: the equation is a simplified, one-compartment form of established glucose-kinetics models (Bergman et al., 1979 [18]; Dalla Man et al., 2007 [19]). The gamma-shaped gut absorption follows Hovorka et al. (2004) [20]; activity energy uses MET values (Ainsworth et al., 2011 [2]); the share of carbohydrate burnt during exercise follows the crossover concept (Brooks and Mercier, 1994 [21]); glycaemic load is GI × available carbohydrate / 100 (Salmerón et al., 1997 [22]). The clearance rates k and the OGTT anchor rises (+3, +5, +8 mmol/L) are approximate values that I chose to be consistent with the WHO (2006) OGTT categories [23]; they are not fitted to measured data.')
remove(r_cite)
# 5 advisor diagram
figure_after(r_adv, '17_advisor_flow.png', 'How a Meal Advisor answer is produced and checked.'); remove(r_adv)
# 6 learning diagram
figure_after(r_learn, '18_learning_flow.png', 'From a food named by a user to a new model, with consent, checks and a gate.'); remove(r_learn)
# 7 detailed system diagram
figure_after(r_sys, '19_system_detail.png', 'Components of the system and how data flows between them.', 6.3); remove(r_sys)
# 8 experiments flow
figure_after(r_exp, '20_experiments_flow.png', 'Each experiment: its input, method, output file and result.', 6.3); remove(r_exp)
# 9 legend for Figure 5: replace the histogram image and extend its caption
cap5 = find(lambda p: p.style.name == 'Caption' and 'Distribution of F1 over the 320 classes' in p.text)
imgp = cap5._p.getprevious()
imgpar = docx.text.paragraph.Paragraph(imgp, cap5._parent)
for r in list(imgpar.runs): r._r.getparent().remove(r._r)
imgpar.add_run().add_picture(CH + '13_per_class_f1.png', width=Inches(5.8))
cap5.add_run(' Orange: the 45 classes with F1 below 50%, which need improvement; teal: the 275 classes at 50% or above. Each bar counts the classes in a 10-point F1 band.')
remove(r_legend)

# corrections
for p in d.paragraphs:
    pass
def replace_in_cells(old, new):
    n = 0
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if old in p.text:
                        full = p.text.replace(old, new)
                        for r in p.runs[1:]: r._r.getparent().remove(r._r)
                        p.runs[0].text = full; n += 1
    return n
assert replace_in_cells('weight decay 1e-4 (0.00001)', 'weight decay 1e-4 (0.0001)')
assert replace_in_cells('I prepared the data and trained the model on Buddy: describe exactly what Lisa wrote or advised, e.g. the cleaning, split, training and evaluation scripts.',
                        'I prepared the data and trained the model on Buddy, using Lisa to write and explain the scripts listed in the next column; I ran them myself.')
assert replace_in_cells('I used SLM is originally proposed. I used it as Meal Advisor (a Chatbot)  running at Buddy’s. I checked AI output by running it rather than by trusting it.',
                        'As proposed, the Meal Advisor (a chatbot) uses this local SLM, running on Buddy. It is part of the product, not a tool used to write it. I checked its output by measuring it (Section 6.4) rather than by trusting it.')
assert replace_in_cells('It provides a textual response from user’s inputted statement. The SLM grounds it’s response from the captured data (week’s log) and does not share outside MealSense.',
                        'It rewords answers that the rule engine drafts from the user\'s own logged data, such as the week\'s meals; a validator checks every number before an answer is shown. Nothing is sent to an outside service.')

# plain-language summaries under technical headings
PLAIN = {
 '2. Tools and techniques': 'one image model finds the foods on a plate and another names them; a food table turns foods into nutrients; a small local text model only words the explanations. Everything runs on one ordinary server, without a paid AI service.',
 '3.3 Data': 'I taught the model with about 300,000 labelled food photos, removed duplicates and photos with contradictory labels, and kept 45,800 photos aside that the model never saw, so that it could be tested fairly.',
 '3.4 Recognition model': 'I started from a model that already recognises everyday objects and retrained it in two steps to recognise 320 foods. On photos it had never seen, its first guess is right about 7 times in 10, and the right food is among its top three guesses almost 9 times in 10.',
 '3.5 Several foods on a plate': 'when a plate holds several foods, the app draws a box around each one and names it separately. When it is unsure, it asks the user instead of guessing.',
 '3.6 Portions, nutrition and glycaemic load': 'the weight of each food comes from the user or from an estimate with a stated margin; nutrients come from the food table; glycaemic load shows how strongly the carbohydrate in the meal is likely to raise blood sugar.',
 '3.7 Glucose model': 'the app estimates the shape of the blood-sugar rise after a meal, how high and for how long, and how a walk or a smaller portion would change it. It is an educational estimate, not a measurement.',
 '3.8 Grounded Meal Advisor': 'the Advisor\'s facts always come from the user\'s own records and the food table. The language model only rephrases them, and any answer containing something that cannot be checked is replaced by the plain rule-based answer.',
 '3.9 Learning safely from users': 'the app remembers a food a user has named straight away. Photos are used to improve the model only with permission and after checks, and a new model must prove it is at least as good before it replaces the current one.',
 '4.4 Implementation and optimisation choices': 'the app is built to be fast and cheap on one server, to keep working when the optional language model is off, and to let an administrator adjust its settings without changing code.',
 '5. Experiments': 'every claim in this report was tested. Each of the eight experiments has fixed test data, a method and a saved result file, so it can be repeated.',
 '6.1 Recognition': 'the model is right first time for about 7 in 10 photos and has the answer in its top three for about 9 in 10. It is weakest on local noodle and rice dishes that look alike. When it says it is unsure, it usually is, so the app asks the user.',
 '6.2 Speed': 'a full plate is analysed in about a quarter of a second. The optional language model adds 1 to 5 seconds, so its text appears after the results.',
 '6.4 Language and text matching': 'the small language model was less accurate than simple rules at reading meal descriptions and matching food names, so it is only used to reword answers, and its output is always checked.',
 '6.6 Glucose model scenarios': 'for this meal, eating a third less rice lowers the estimated blood-sugar peak about ten times more than a 15-minute walk.',
 '7.2 Trade-offs and threats to validity': 'the results come from curated test photos. Photos taken at a hawker stall may give lower accuracy, which the planned user trial will measure.',
}
done = set()
for p in list(d.paragraphs):
    if p.style.name.startswith('Heading') and p.text.strip() in PLAIN and p.text.strip() not in done:
        plain_after(p, PLAIN[p.text.strip()]); done.add(p.text.strip())
missing = set(PLAIN) - done
print('plain added', len(done), 'missing', missing)

# references 18-23
refs = [p for p in d.paragraphs if p.text.startswith('Health Promotion Board, Singapore. Energy & Nutrient Composition')]
last = refs[0]
NEW_REFS = [
 'Bergman RN, Ider YZ, Bowden CR, Cobelli C. Quantitative estimation of insulin sensitivity. Am J Physiol. 1979;236(6):E667–E677.',
 'Dalla Man C, Rizza RA, Cobelli C. Meal simulation model of the glucose-insulin system. IEEE Trans Biomed Eng. 2007;54(10):1740–1749.',
 'Hovorka R, et al. Nonlinear model predictive control of glucose concentration in subjects with type 1 diabetes. Physiol Meas. 2004;25(4):905–920.',
 'Brooks GA, Mercier J. Balance of carbohydrate and lipid utilization during exercise: the "crossover" concept. J Appl Physiol. 1994;76(6):2253–2261.',
 'Salmerón J, et al. Dietary fiber, glycemic load, and risk of non-insulin-dependent diabetes mellitus in women. JAMA. 1997;277(6):472–477.',
 'World Health Organization. Definition and diagnosis of diabetes mellitus and intermediate hyperglycaemia. Geneva: WHO; 2006.',
]
for t in NEW_REFS:
    el = copy.deepcopy(last._p); last._p.addnext(el)
    np = docx.text.paragraph.Paragraph(el, last._parent)
    for r in np.runs[1:]: r._r.getparent().remove(r._r)
    np.runs[0].text = t; last = np

# renumber figures in document order and update in-text references
caps = [p for p in d.paragraphs if p.style.name == 'Caption' and re.match(r'Figure \d+\.', p.text)]
mapping = {}
for i, p in enumerate(caps, 1):
    old = int(re.match(r'Figure (\d+)\.', p.text).group(1))
    if old != 0: mapping[old] = i
    p.runs[0].text = f'Figure {i}. '
print('figures', len(caps), mapping)
def fix_refs(p):
    for r in p.runs:
        if 'Figure' in r.text:
            r.text = re.sub(r'Figure (\d+)(?!\.\s)', lambda m: f'Figure {mapping.get(int(m.group(1)), int(m.group(1)))}', r.text)
for p in d.paragraphs:
    if p.style.name != 'Caption': fix_refs(p)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: fix_refs(p)
d.save('MealSenseAI-Final-Report.docx'); print('saved')
